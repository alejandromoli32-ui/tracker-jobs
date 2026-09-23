"""Excel (.xlsx) job tracker generator using openpyxl.

Produces highly styled, interactive workbooks with clickable hyperlinks,
conditional color fills based on affinity scores, status validation dropdowns,
and frozen header rows.
"""

from __future__ import annotations

import os
import re
import zipfile
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Union

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from job_hunter.models import TrackerEntry

CONTROL_CHARS_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
DATE_DETECTION_REGEX = re.compile(r"\bfecha\s*(?:de\s*)?detecci[oó]n\b", re.IGNORECASE)


def _clean_cell_value(val: Any) -> Any:
    """Strip illegal ASCII control characters to prevent openpyxl IllegalCharacterError."""
    if isinstance(val, str):
        return CONTROL_CHARS_RE.sub("", val)
    return val


DEFAULT_HEADERS: List[str] = [
    "ID/Hash",
    "Fecha Detección",
    "Título",
    "Empresa",
    "Score %",
    "Modalidad",
    "Enlace Directo",
    "Salario",
    "Resumen de Afinidad",
    "Estado",
    "Notas Usuario",
]

EXCEL_STATUS_OPTIONS: List[str] = [
    "Nueva",
    "Por revisar",
    "Postulado",
    "Descartado",
    "Entrevista",
]

# Color palette specifications
NAVY_FILL_HEX = "1B365D"
BORDER_GRAY_HEX = "D3D3D3"
LINK_BLUE_HEX = "0563C1"

# Score fill palettes
GREEN_FILL_HEX = "D9EAD3"
GREEN_FONT_HEX = "274E13"

YELLOW_FILL_HEX = "FFF2CC"
YELLOW_FONT_HEX = "7F6000"

RED_FILL_HEX = "F4CCCC"
RED_FONT_HEX = "783F04"

# Column minimum widths
COLUMN_MIN_WIDTHS: Dict[int, int] = {
    1: 16,  # ID/Hash
    2: 18,  # Fecha Detección
    3: 38,  # Título
    4: 26,  # Empresa
    5: 12,  # Score %
    6: 18,  # Modalidad
    7: 16,  # Enlace Directo
    8: 22,  # Salario
    9: 46,  # Resumen de Afinidad
    10: 16, # Estado
    11: 32, # Notas Usuario
}


class ExcelTracker:
    """Manages creation, styling, and reading of openpyxl Excel job tracker files."""

    def __init__(
        self,
        output_path: Optional[Union[str, Path]] = None,
        sheet_name: str = "Vacantes_Vias_SST",
    ) -> None:
        self.output_path = Path(output_path) if output_path else None
        self.sheet_name = sheet_name

    @classmethod
    def save(
        cls,
        entries: Sequence[TrackerEntry],
        output_path: Optional[Union[str, Path]] = None,
        preserve_existing: bool = True,
        sheet_name: str = "Vacantes_Vias_SST",
    ) -> Path:
        """Save tracker entries to an Excel workbook (.xlsx).

        If preserve_existing is True and output_path exists, an incremental merge
        is performed prior to rewriting the workbook.
        """
        if output_path is None:
            raise ValueError("output_path must be provided to ExcelTracker.save()")

        path = Path(output_path).resolve()
        path.parent.mkdir(parents=True, exist_ok=True)

        merged_entries: List[TrackerEntry] = list(entries)
        if preserve_existing and path.exists() and path.stat().st_size > 0:
            from job_hunter.storage.incremental import IncrementalTracker
            merged_entries = IncrementalTracker.merge(path, entries)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_name

        # 1. Write and style header row
        ws.append(DEFAULT_HEADERS)
        ws.row_dimensions[1].height = 28.0

        header_fill = PatternFill(start_color=NAVY_FILL_HEX, end_color=NAVY_FILL_HEX, fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

        thin_side = Side(style="thin", color=BORDER_GRAY_HEX)
        cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

        for col_idx in range(1, len(DEFAULT_HEADERS) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_align
            cell.border = cell_border

        # 2. Write data rows
        regular_font = Font(name="Calibri", size=10)
        link_font = Font(name="Calibri", size=10, color=LINK_BLUE_HEX, underline="single")

        for row_idx, entry in enumerate(merged_entries, start=2):
            ws.row_dimensions[row_idx].height = 22.0

            # Compute score decimal for Excel percentage display (e.g. 0.88 for 88%)
            raw_score = float(entry.score) if entry.score is not None else 0.0
            excel_score_val = raw_score / 100.0 if raw_score > 1.0 else raw_score

            # Clickable formula
            direct_url = (entry.direct_url or "").strip()
            clean_url = _clean_cell_value(direct_url)
            safe_url = clean_url.replace('"', "%22")
            hyperlink_formula = f'=HYPERLINK("{safe_url}", "Ver Vacante")' if safe_url else ""

            row_data = [
                _clean_cell_value(entry.id),
                _clean_cell_value(entry.fecha_deteccion or ""),
                _clean_cell_value(entry.title or ""),
                _clean_cell_value(entry.company or ""),
                excel_score_val,
                _clean_cell_value(entry.modality.capitalize() if entry.modality else "Remoto"),
                hyperlink_formula,
                _clean_cell_value(entry.salary_range or ""),
                _clean_cell_value(entry.explanatory_summary or ""),
                _clean_cell_value(entry.estado or "Nueva"),
                _clean_cell_value(entry.notas_usuario or ""),
            ]
            ws.append(row_data)

            # Apply cell-specific alignments, borders, and styles
            # Col 1: ID
            c_id = ws.cell(row=row_idx, column=1)
            c_id.alignment = Alignment(horizontal="center", vertical="center")
            c_id.font = regular_font
            c_id.border = cell_border

            # Col 2: Fecha Detección
            c_date = ws.cell(row=row_idx, column=2)
            c_date.alignment = Alignment(horizontal="center", vertical="center")
            c_date.font = regular_font
            c_date.border = cell_border

            # Col 3: Título
            c_title = ws.cell(row=row_idx, column=3)
            c_title.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            c_title.font = regular_font
            c_title.border = cell_border

            # Col 4: Empresa
            c_comp = ws.cell(row=row_idx, column=4)
            c_comp.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            c_comp.font = regular_font
            c_comp.border = cell_border

            # Col 5: Score %
            c_score = ws.cell(row=row_idx, column=5)
            c_score.number_format = "0.0%"
            c_score.alignment = Alignment(horizontal="center", vertical="center")
            c_score.border = cell_border

            pct = excel_score_val * 100.0
            if pct >= 70.0:
                c_score.fill = PatternFill(start_color=GREEN_FILL_HEX, end_color=GREEN_FILL_HEX, fill_type="solid")
                c_score.font = Font(name="Calibri", size=10, bold=True, color=GREEN_FONT_HEX)
            elif pct >= 50.0:
                c_score.fill = PatternFill(start_color=YELLOW_FILL_HEX, end_color=YELLOW_FILL_HEX, fill_type="solid")
                c_score.font = Font(name="Calibri", size=10, bold=True, color=YELLOW_FONT_HEX)
            else:
                c_score.fill = PatternFill(start_color=RED_FILL_HEX, end_color=RED_FILL_HEX, fill_type="solid")
                c_score.font = Font(name="Calibri", size=10, bold=True, color=RED_FONT_HEX)

            # Col 6: Modalidad
            c_mod = ws.cell(row=row_idx, column=6)
            c_mod.alignment = Alignment(horizontal="center", vertical="center")
            c_mod.font = regular_font
            c_mod.border = cell_border

            # Col 7: Enlace Directo
            c_link = ws.cell(row=row_idx, column=7)
            c_link.alignment = Alignment(horizontal="center", vertical="center")
            c_link.font = link_font
            c_link.border = cell_border

            # Col 8: Salario
            c_sal = ws.cell(row=row_idx, column=8)
            c_sal.alignment = Alignment(horizontal="left", vertical="center")
            c_sal.font = regular_font
            c_sal.border = cell_border

            # Col 9: Resumen de Afinidad
            c_res = ws.cell(row=row_idx, column=9)
            c_res.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            c_res.font = regular_font
            c_res.border = cell_border

            # Col 10: Estado
            c_est = ws.cell(row=row_idx, column=10)
            c_est.alignment = Alignment(horizontal="center", vertical="center")
            c_est.font = regular_font
            c_est.border = cell_border

            # Col 11: Notas Usuario
            c_not = ws.cell(row=row_idx, column=11)
            c_not.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            c_not.font = regular_font
            c_not.border = cell_border

        # 3. Data validation for Estado (Column J)
        dv_formula = f'"{",".join(EXCEL_STATUS_OPTIONS)}"'
        dv = DataValidation(
            type="list",
            formula1=dv_formula,
            allow_blank=True,
            showDropDown=False,
        )
        ws.add_data_validation(dv)
        max_dv_row = max(len(merged_entries) + 200, 500)
        dv.add(f"J2:J{max_dv_row}")

        # 4. Freeze top header row
        ws.freeze_panes = "A2"

        # 5. Auto-fit column widths with padding
        for col_idx in range(1, len(DEFAULT_HEADERS) + 1):
            col_letter = get_column_letter(col_idx)
            header_len = len(DEFAULT_HEADERS[col_idx - 1])
            max_len = header_len

            for row_idx in range(2, len(merged_entries) + 2):
                val = ws.cell(row=row_idx, column=col_idx).value
                if val is not None:
                    s_val = str(val)
                    if s_val.startswith("=HYPERLINK"):
                        s_val = "Ver Vacante"
                    # Take line length for multiline texts
                    lines = s_val.splitlines()
                    val_len = max(len(line) for line in lines) if lines else 0
                    if val_len > max_len:
                        max_len = val_len

            min_w = COLUMN_MIN_WIDTHS.get(col_idx, 12)
            calculated_w = max(max_len + 4, min_w)
            # Enforce reasonable ceiling per column to avoid excessively wide columns
            if col_idx in (3, 9, 11):  # Title, Summary, Notes
                calculated_w = min(calculated_w, 65)
            else:
                calculated_w = min(calculated_w, 35)

            ws.column_dimensions[col_letter].width = calculated_w

        # Save workbook
        wb.save(path)
        wb.close()
        return path

    def write(self, entries: Sequence[TrackerEntry], preserve_existing: bool = True) -> Path:
        """Instance helper to save using instance configuration."""
        if not self.output_path:
            raise ValueError("output_path was not set on ExcelTracker instance")
        return self.save(
            entries=entries,
            output_path=self.output_path,
            preserve_existing=preserve_existing,
            sheet_name=self.sheet_name,
        )

    @classmethod
    def load(cls, file_path: Union[str, Path]) -> List[TrackerEntry]:
        """Read an existing Excel tracker workbook into a list of TrackerEntry objects."""
        path = Path(file_path).resolve()
        if not path.exists() or path.stat().st_size == 0:
            return []

        try:
            wb = openpyxl.load_workbook(path, data_only=False)
        except Exception:
            return []

        ws = wb.active

        # Read header row
        header_row = [str(c).strip() if c is not None else "" for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True), [])]
        if not header_row:
            wb.close()
            return []

        # Build column index mapping
        col_map: Dict[str, int] = {}
        for idx, h in enumerate(header_row):
            h_lower = h.lower().strip()
            if re.search(r"\b(id|hash)\b", h_lower):
                col_map["id"] = idx
            elif DATE_DETECTION_REGEX.search(h_lower):
                col_map["fecha"] = idx
            elif "título" in h_lower or "titulo" in h_lower:
                col_map["title"] = idx
            elif "empresa" in h_lower or "company" in h_lower:
                col_map["company"] = idx
            elif "score" in h_lower or "puntaje" in h_lower:
                col_map["score"] = idx
            elif "modalidad" in h_lower or "modality" in h_lower:
                col_map["modality"] = idx
            elif "enlace" in h_lower or "url" in h_lower or "link" in h_lower:
                col_map["url"] = idx
            elif "salario" in h_lower or "salary" in h_lower:
                col_map["salary"] = idx
            elif "resumen" in h_lower or "afinidad" in h_lower or "summary" in h_lower:
                col_map["summary"] = idx
            elif "estado" in h_lower or "status" in h_lower:
                col_map["estado"] = idx
            elif "nota" in h_lower or "note" in h_lower:
                col_map["notas"] = idx
            elif "fecha" in h_lower and "postula" not in h_lower and "fecha" not in col_map:
                col_map["fecha"] = idx

        entries: List[TrackerEntry] = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not any(row):
                continue

            def _get(field_key: str, default: Any = "") -> Any:
                idx = col_map.get(field_key)
                if idx is not None and idx < len(row):
                    val = row[idx]
                    return val if val is not None else default
                return default

            v_id = str(_get("id", "")).strip()
            if not v_id:
                continue

            # Extract clean URL from =HYPERLINK formula or text
            raw_url = str(_get("url", "")).strip()
            match = re.search(r'=HYPERLINK\("([^"]+)"', raw_url)
            clean_url_val = match.group(1) if match else raw_url

            # Parse score
            raw_score_val = _get("score", 0.0)
            score_num = 0.0
            if isinstance(raw_score_val, (int, float)):
                score_num = float(raw_score_val)
                # Convert 0.88 -> 88.0
                if 0.0 < score_num <= 1.0:
                    score_num = round(score_num * 100.0, 1)
            elif isinstance(raw_score_val, str):
                cleaned_sc = re.sub(r"[^\d.]", "", raw_score_val)
                try:
                    score_num = float(cleaned_sc)
                    if 0.0 < score_num <= 1.0:
                        score_num = round(score_num * 100.0, 1)
                except ValueError:
                    score_num = 0.0

            # Clamp score to [0.0, 100.0]
            score_num = max(0.0, min(100.0, score_num))

            title_val = str(_get("title", "Vacante")).strip()
            company_val = str(_get("company", "Confidencial")).strip()
            modality_val = str(_get("modality", "remoto")).strip().lower()
            salary_val = str(_get("salary", "")).strip() or None
            summary_val = str(_get("summary", "")).strip()
            estado_val = str(_get("estado", "Nueva")).strip() or "Nueva"
            fecha_val = str(_get("fecha", "")).strip()
            notas_val = str(_get("notas", "")).strip()

            entry = TrackerEntry(
                id=v_id,
                title=title_val,
                company=company_val,
                direct_url=clean_url_val,
                full_description=summary_val,
                location="Colombia",
                modality=modality_val,
                salary_range=salary_val,
                source_portal="excel",
                score=score_num,
                explanatory_summary=summary_val,
                estado=estado_val,
                fecha_deteccion=fecha_val,
                notas_usuario=notas_val,
            )
            entries.append(entry)

        wb.close()
        return entries
