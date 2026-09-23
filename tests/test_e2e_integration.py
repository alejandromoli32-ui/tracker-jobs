"""Comprehensive End-to-End (E2E) Integration Test Suite.

Validates:
1. Full pipeline execution via Python API (`run_pipeline`) and CLI runner (`main`).
2. Excel output aesthetics, structure, `=HYPERLINK`, color fills, dropdown validation, freeze panes.
3. Incremental state preservation across recurring runs (user status, notes, dates).
4. Multi-profile independence (Profiles 1-4) and zero-code extensibility (Profile 5).
5. All Acceptance Criteria from ORIGINAL_REQUEST.md (lines 36-52):
   - Functional extraction and data normalization.
   - Network error resilience and error domain isolation.
   - Presential/obra vacancies filtered or penalized.
   - Score > 70 matches Civil/Vías/SST profiles (9 Canonical Benchmark Jobs).
   - Multi-profile requires no code modifications.
   - Styled Excel and incremental state preservation.
6. Challenger M4 hardening mitigations (control characters, corrupt files, first-run dedup, score clamping, date regex).
"""

from __future__ import annotations

import csv
import io
import json
import re
import sys
from pathlib import Path
from typing import Any, Dict, List
from unittest.mock import MagicMock, patch

import openpyxl
import pytest
from openpyxl.styles import PatternFill

from job_hunter.cli import build_parser, main, run_pipeline
from job_hunter.config import list_available_profiles, load_profile
from job_hunter.dedup import Deduplicator
from job_hunter.matcher import AffinityScorer
from job_hunter.models import NormalizedVacancy, ProfileConfig, TrackerEntry
from job_hunter.remote_verifier import RemoteVerifier
from job_hunter.scrapers.base import BaseScraper
from job_hunter.scrapers.engine import JobAggregationEngine
from job_hunter.storage.csv_exporter import CSVTracker
from job_hunter.storage.excel import (
    DEFAULT_HEADERS,
    GREEN_FILL_HEX,
    GREEN_FONT_HEX,
    RED_FILL_HEX,
    YELLOW_FILL_HEX,
    ExcelTracker,
)
from job_hunter.storage.incremental import IncrementalTracker


# =============================================================================
# FIXTURES & CANONICAL BENCHMARK DATASET
# =============================================================================

def _make_canon(
    title: str,
    company: str,
    url: str,
    desc: str,
    location: str,
    modality: str,
    salary: str,
    portal: str,
    pub_date: str,
    category: str,
) -> NormalizedVacancy:
    return NormalizedVacancy(
        id=Deduplicator.compute_id(title, company, url),
        title=title,
        company=company,
        direct_url=url,
        full_description=desc,
        location=location,
        modality=modality,
        salary_range=salary,
        source_portal=portal,
        publication_date=pub_date,
        extra_metadata={"category": category},
    )


@pytest.fixture
def canonical_benchmark_vacancies() -> List[NormalizedVacancy]:
    """The 9 Canonical Benchmark Job Vacancies specified in ORIGINAL_REQUEST.md / TEST_INFRA.md:
    - 3 Ideal Remote Vías + SST (Expected Score >= 70, ACCEPT)
    - 3 On-Site Obra / Campamento (Expected Score <= 20, DISQUALIFIED / PENALIZED)
    - 3 Non-relevant Civil Engineering (Expected Score < 50, REJECT / FILTERED)
    """
    return [
        # --- 3 Ideal Remote Vías + SST Vacancies ---
        _make_canon(
            title="Directora de Interventoría Vías y SST Remota",
            company="Concesiones Viales Andinas S.A.S.",
            url="https://computrabajo.com/co/job/v01",
            desc=(
                "Firma de infraestructura requiere Ingeniera Civil con más de 16 años de trayectoria general. "
                "Especialista en Vías Terrestres con más de 6 años de experiencia en interventoría de carreteras y pavimentos. "
                "Especialización en Seguridad y Salud en el Trabajo (SST) con Licencia SST vigente y tarjeta COPNIA. "
                "Modalidad 100% teletrabajo / virtual desde Colombia para auditoría técnica de contratos y SG-SST."
            ),
            location="Colombia (Remoto)",
            modality="remoto",
            salary="$12.000.000 - $16.000.000 COP",
            portal="computrabajo",
            pub_date="2026-09-20",
            category="ideal",
        ),
        _make_canon(
            title="Coordinadora SST e Infraestructura Vial Virtual",
            company="Infraestructura & Consultoría SAS",
            url="https://elempleo.com/co/job/v02",
            desc=(
                "Consultora requiere profesional en Ingeniería Civil con tarjeta COPNIA. "
                "Experiencia mayor a 15 años en gerencia de proyectos viales. "
                "Especialista en Vías y Especialista en SST con resolución y licencia activa. "
                "100% trabajo remoto / home office para elaboración de PESV y revisión de diseños geométricos."
            ),
            location="Bogotá (Remoto)",
            modality="remoto",
            salary="$11.000.000 COP",
            portal="elempleo",
            pub_date="2026-09-19",
            category="ideal",
        ),
        _make_canon(
            title="Asesora Técnica Especialista en Pavimentos y HSE",
            company="Consorcio Vial Nacional",
            url="https://linkedin.com/jobs/view/v03",
            desc=(
                "Importante consorcio vial solicita Ingeniero Civil Senior con 15+ años de experiencia. "
                "Especialización en Vías con énfasis en diseño de pavimentos asfálticos y rígidos (INVIAS). "
                "Especialización en Salud Ocupacional / SST con licencia. "
                "Modalidad virtual / teletrabajo nacional para asesoría normativa y seguimiento de indicadores SST."
            ),
            location="Medellín (Teletrabajo)",
            modality="remoto",
            salary="$10.000.000 - $14.000.000 COP",
            portal="linkedin",
            pub_date="2026-09-18",
            category="ideal",
        ),

        # --- 3 On-Site Obra / Campamento Presential Vacancies (Must be Disqualified/Penalized) ---
        _make_canon(
            title="Ingeniero Residente de Pavimentación en Obra",
            company="Constructora del Sol S.A.",
            url="https://computrabajo.com/co/job/v04",
            desc=(
                "Se requiere Ingeniero Civil con 8 años de experiencia en pavimentación asfáltica. "
                "Disponibilidad 100% presencial en obra, campamento 21x7 en frente de obra en Puerto Boyacá. "
                "Control de maquinaria y personal en terreno. Abstenerse personas para trabajo remoto."
            ),
            location="Puerto Boyacá",
            modality="presencial",
            salary="$7.000.000 COP",
            portal="computrabajo",
            pub_date="2026-09-17",
            category="presential",
        ),
        _make_canon(
            title="Inspector SST en Frente de Obra 4G",
            company="Vías de la Montaña SAS",
            url="https://elempleo.com/co/job/v05",
            desc=(
                "Se solicita Inspector de Seguridad y Salud en el Trabajo. "
                "100% presencial en frente de obra vial concesión 4G. "
                "Turnos rotativos en sitio de excavación y túneles. Vivir en campamento obligatorio."
            ),
            location="Dabeiba, Antioquia",
            modality="presencial",
            salary="$5.500.000 COP",
            portal="elempleo",
            pub_date="2026-09-16",
            category="presential",
        ),
        _make_canon(
            title="Director de Obra Presencial Túneles y Vías",
            company="Megaobras Colombia",
            url="https://linkedin.com/jobs/view/v06",
            desc=(
                "Director de Obra presencial para construcción de intercambiador vial. "
                "Trabajo de campo en terreno 100% presencial. Supervisión directa en el frente de obra."
            ),
            location="Cali, Valle del Cauca",
            modality="presencial",
            salary="$14.000.000 COP",
            portal="linkedin",
            pub_date="2026-09-15",
            category="presential",
        ),

        # --- 3 Non-Relevant Civil Engineering Vacancies (Score < 50, Must be Filtered) ---
        _make_canon(
            title="Ingeniero Calculista Estructural Edificaciones",
            company="Diseños & Estructuras Urbanas",
            url="https://computrabajo.com/co/job/v07",
            desc=(
                "Diseño estructural de edificios residenciales en concreto reforzado bajo NSR-10. "
                "Manejo avanzado de ETABS, SAP2000 y SAFE. Modalidad remota / teletrabajo. "
                "Elaboración de memorias de cálculo y planos estructurales para curadurías urbanas."
            ),
            location="Bogotá (Remoto)",
            modality="remoto",
            salary="$6.000.000 COP",
            portal="computrabajo",
            pub_date="2026-09-14",
            category="nonrelevant",
        ),
        _make_canon(
            title="Residente Hidráulico y Redes Sanitarias",
            company="Aguas & Proyectos Colombia",
            url="https://elempleo.com/co/job/v08",
            desc=(
                "Modelación hidráulica de acueductos y alcantarillados usando WaterCAD y SewerCAD. "
                "Elaboración de planos hidrosanitarios y redes de aguas lluvias. Modalidad remota para consultoría."
            ),
            location="Remoto",
            modality="remoto",
            salary="$5.500.000 COP",
            portal="elempleo",
            pub_date="2026-09-13",
            category="nonrelevant",
        ),
        _make_canon(
            title="Ingeniero Civil Geotécnico de Fundaciones",
            company="Geotecnia Avanzada Ltda",
            url="https://linkedin.com/jobs/view/v09",
            desc=(
                "Análisis geotécnico de capacidad portante y asentamientos para cimentaciones profundas en edificaciones. "
                "Uso de Plaxis y Slide. Trabajo virtual / remoto."
            ),
            location="Remoto",
            modality="remoto",
            salary="$6.500.000 COP",
            portal="linkedin",
            pub_date="2026-09-12",
            category="nonrelevant",
        ),
    ]



class MockCustomScraper(BaseScraper):
    """Deterministic mock scraper returning provided vacancies."""
    def __init__(self, vacancies: List[NormalizedVacancy], name: str = "mock_portal", client=None):
        super().__init__(client=client)
        self.name = name
        self.preset_vacancies = vacancies

    def search(self, query: str = "", max_pages: int = 1) -> List[NormalizedVacancy]:
        return list(self.preset_vacancies)



# =============================================================================
# E2E TEST 1: PIPELINE EXECUTION VIA PYTHON API & CLI RUNNER
# =============================================================================

class TestE2EPipelineExecution:
    """Test full pipeline execution through both programmatic API and CLI runner."""

    def test_run_pipeline_programmatic_api(self, tmp_path: Path):
        """Execute run_pipeline() directly with Profile 1 in offline mode."""
        res = run_pipeline(
            profile_input="profile_1_civil_vias_sst",
            output_dir=tmp_path / "api_out",
            export_format="both",
            offline=True,
            min_score=70.0,
            quiet=True,
        )

        assert res["profile_id"] == "profile_1_civil_vias_sst"
        assert res["total_scraped"] > 0
        assert res["remote_verified"] > 0
        assert res["scored_gte_70"] > 0
        assert res["exported_count"] == res["scored_gte_70"]
        assert len(res["exported_files"]) == 2

        xlsx_path = Path(res["exported_files"][0])
        csv_path = Path(res["exported_files"][1])
        assert xlsx_path.exists() and xlsx_path.stat().st_size > 0
        assert csv_path.exists() and csv_path.stat().st_size > 0

    def test_cli_runner_main_with_arguments(self, tmp_path: Path):
        """Execute main() via argv simulation for Profile 1."""
        out_dir = tmp_path / "cli_out"
        argv = [
            "--profile", "profile_1_civil_vias_sst",
            "--offline",
            "--output-dir", str(out_dir),
            "--format", "both",
            "--min-score", "70.0",
            "--quiet",
        ]

        exit_code = main(argv)
        assert exit_code == 0

        expected_xlsx = out_dir / "profile_1_civil_vias_sst_tracker.xlsx"
        expected_csv = out_dir / "profile_1_civil_vias_sst_tracker.csv"
        assert expected_xlsx.exists()
        assert expected_csv.exists()

    def test_cli_list_profiles_flag(self, capsys):
        """Verify --list-profiles lists all configured profiles and exits cleanly."""
        exit_code = main(["--list-profiles"])
        assert exit_code == 0

        captured = capsys.readouterr().out
        assert "AVAILABLE PROFILES" in captured
        assert "profile_1_civil_vias_sst" in captured
        assert "profile_2_gerente_proyectos_viales" in captured
        assert "profile_3_modelador_bim_vias" in captured
        assert "profile_4_consultor_hseq_sig" in captured

    def test_cli_error_handling_invalid_profile(self, capsys):
        """Verify CLI returns exit code 1 with clear error when given a nonexistent profile."""
        exit_code = main(["--profile", "non_existent_profile_xyz", "--offline"])
        assert exit_code == 1

    def test_cli_parser_defaults_and_options(self):
        """Verify build_parser() configuration, defaults, and option flags."""
        parser = build_parser()
        args = parser.parse_args([])
        assert args.profile == "profile_1_civil_vias_sst"
        assert args.portals == "all"
        assert args.format == "both"
        assert args.offline is False
        assert args.max_pages == 2
        assert args.output_dir == "output"


# =============================================================================
# E2E TEST 2: EXCEL WORKBOOK STYLING, HYPERLINKS, AND STRUCTURE
# =============================================================================

class TestE2EExcelOutputVerification:
    """Verify generated Excel file complies with all styling and structure requirements."""

    def test_excel_file_complete_structure_and_styling(self, tmp_path: Path):
        out_dir = tmp_path / "excel_test"
        run_pipeline(
            profile_input="profile_1_civil_vias_sst",
            output_dir=out_dir,
            export_format="xlsx",
            offline=True,
            min_score=50.0,
            quiet=True,
        )

        excel_path = out_dir / "profile_1_civil_vias_sst_tracker.xlsx"
        assert excel_path.exists()

        wb = openpyxl.load_workbook(excel_path)
        ws = wb.active

        # 1. Verify headers
        header_vals = [ws.cell(row=1, column=c).value for c in range(1, len(DEFAULT_HEADERS) + 1)]
        assert header_vals == DEFAULT_HEADERS

        # 2. Verify freeze panes
        assert ws.freeze_panes == "A2"

        # 3. Verify data rows exist
        assert ws.max_row >= 2

        # 4. Verify clickable hyperlinks in column 7 (Enlace Directo)
        for row_idx in range(2, ws.max_row + 1):
            link_cell = ws.cell(row=row_idx, column=7)
            val = str(link_cell.value or "")
            if val:
                assert val.startswith('=HYPERLINK("'), f"Row {row_idx} missing =HYPERLINK formula: {val}"
                assert "Ver Vacante" in val

        # 5. Verify Score formatting and conditional colors in column 5
        for row_idx in range(2, ws.max_row + 1):
            score_cell = ws.cell(row=row_idx, column=5)
            assert score_cell.number_format == "0.0%"
            fill_color = score_cell.fill.start_color.rgb if score_cell.fill else None
            assert fill_color is not None, f"Row {row_idx} has no fill color"
            fill_hex = str(fill_color).upper()
            score_num = float(score_cell.value or 0.0) * 100.0

            if score_num >= 70.0:
                assert fill_hex.endswith(GREEN_FILL_HEX)
                assert score_cell.font.bold is True
                assert score_cell.font.color.rgb.upper().endswith(GREEN_FONT_HEX)
            elif score_num >= 50.0:
                assert fill_hex.endswith(YELLOW_FILL_HEX)
            else:
                assert fill_hex.endswith(RED_FILL_HEX)

        # 6. Verify status dropdown validation on column 10 (Estado)
        assert len(ws.data_validations.dataValidation) > 0
        dv = ws.data_validations.dataValidation[0]
        assert dv.type == "list"
        assert "Nueva" in dv.formula1
        assert "Postulado" in dv.formula1
        assert "Descartado" in dv.formula1

        # 7. Verify column widths are configured
        for col_idx in range(1, len(DEFAULT_HEADERS) + 1):
            col_letter = openpyxl.utils.get_column_letter(col_idx)
            w = ws.column_dimensions[col_letter].width
            assert w is not None and w >= 12

        wb.close()


# =============================================================================
# E2E TEST 3: INCREMENTAL STATE PRESERVATION
# =============================================================================

class TestE2EIncrementalStatePreservation:
    """Verify that recurring pipeline runs preserve user-assigned candidate state."""

    def test_incremental_workflow_preserves_postulado_and_notes(self, tmp_path: Path):
        out_dir = tmp_path / "incremental_e2e"
        excel_path = out_dir / "profile_1_civil_vias_sst_tracker.xlsx"
        csv_path = out_dir / "profile_1_civil_vias_sst_tracker.csv"

        # Step 1: Initial pipeline run
        run_pipeline(
            profile_input="profile_1_civil_vias_sst",
            output_dir=out_dir,
            export_format="both",
            offline=True,
            min_score=70.0,
            quiet=True,
        )

        entries_initial = ExcelTracker.load(excel_path)
        assert len(entries_initial) >= 2
        target_id_1 = entries_initial[0].id
        target_id_2 = entries_initial[1].id
        initial_date_1 = entries_initial[0].fecha_deteccion

        # Step 2: User manually updates Excel tracker
        wb = openpyxl.load_workbook(excel_path)
        ws = wb.active
        for r in range(2, ws.max_row + 1):
            c_id = str(ws.cell(row=r, column=1).value)
            if c_id == target_id_1:
                ws.cell(row=r, column=10).value = "Postulado"
                ws.cell(row=r, column=11).value = "CV enviado por Computrabajo con carta de motivación"
            elif c_id == target_id_2:
                ws.cell(row=r, column=10).value = "Descartado"
                ws.cell(row=r, column=11).value = "Descartado por salario fuera de expectativa"
        wb.save(excel_path)
        wb.close()

        # Step 3: Run pipeline second time (simulating recurring hunt)
        run_pipeline(
            profile_input="profile_1_civil_vias_sst",
            output_dir=out_dir,
            export_format="xlsx",
            offline=True,
            min_score=70.0,
            quiet=True,
        )

        # Step 4: Validate state was 100% preserved
        entries_second_run = ExcelTracker.load(excel_path)
        by_id = {e.id: e for e in entries_second_run}

        assert target_id_1 in by_id
        entry_1 = by_id[target_id_1]
        assert entry_1.estado == "Postulado"
        assert entry_1.notas_usuario == "CV enviado por Computrabajo con carta de motivación"
        assert entry_1.fecha_deteccion == initial_date_1

        assert target_id_2 in by_id
        entry_2 = by_id[target_id_2]
        assert entry_2.estado == "Descartado"
        assert entry_2.notas_usuario == "Descartado por salario fuera de expectativa"


# =============================================================================
# E2E TEST 4: MULTI-PROFILE INDEPENDENCE & ZERO-CODE EXTENSIBILITY
# =============================================================================

class TestE2EMultiProfileIndependence:
    """Verify independent multi-profile execution and zero-code extensibility."""

    def test_four_profiles_produce_dedicated_trackers(self, tmp_path: Path):
        """Execute all 4 standard profiles and verify dedicated tracker files."""
        out_dir = tmp_path / "multi_profiles"
        profiles = [
            "profile_1_civil_vias_sst",
            "profile_2_gerente_proyectos_viales",
            "profile_3_modelador_bim_vias",
            "profile_4_consultor_hseq_sig",
        ]

        for p_id in profiles:
            res = run_pipeline(
                profile_input=p_id,
                output_dir=out_dir,
                export_format="both",
                offline=True,
                quiet=True,
            )
            assert res["profile_id"] == p_id
            xlsx_f = out_dir / f"{p_id}_tracker.xlsx"
            csv_f = out_dir / f"{p_id}_tracker.csv"
            assert xlsx_f.exists()
            assert csv_f.exists()

        # Verify all 8 files exist without file overwriting collisions
        all_files = list(out_dir.glob("*"))
        assert len(all_files) == 8

    def test_dynamic_fifth_profile_zero_code_modification(self, tmp_path: Path):
        """Acceptance Criterion line 47: Adding a new profile requires ZERO code modifications.
        Create a 5th profile dynamically in YAML and execute via CLI runner.
        """
        profiles_dir = tmp_path / "custom_profiles"
        profiles_dir.mkdir()
        profile_5_yaml = profiles_dir / "profile_5_especialista_geotecnia_vial.yaml"
        profile_5_content = """
id: "profile_5_especialista_geotecnia_vial"
name: "Especialista en Geotecnia Vial & Estabilidad de Taludes"
target_role: "Especialista Geotécnico Vial Remoto"
version: "1.0.0"
active: true
min_total_experience_years: 10

experience:
  min_total_years: 10
  seniority_level: "Senior"
  primary_discipline: "Ingeniería Civil"

specializations:
  geotecnia:
    name: "Geotecnia y Taludes"
    min_years: 5
    weight: 25.0
    keywords: ["geotecnia", "taludes", "estabilidad", "cimentaciones", "suelos"]
  vias:
    name: "Infraestructura Vial"
    min_years: 4
    weight: 25.0
    keywords: ["vías", "pavimentos", "carreteras", "invias"]

modality:
  strictly_remote: true
  allowed_modalities: ["remoto", "virtual", "teletrabajo"]
  disqualifiers: ["100% presencial", "campamento", "frente de obra"]

search_queries:
  - "geotecnia vial remoto"
  - "ingeniero geotecnico virtual"

scoring_weights:
  modality: 25.0
  seniority: 25.0
  vias: 25.0
  sst: 25.0
  synergy_bonus: 10.0
  min_score_to_save: 50.0
"""
        profile_5_yaml.write_text(profile_5_content, encoding="utf-8")

        # Execute pipeline passing file path directly to CLI main()
        out_dir = tmp_path / "out_p5"
        argv = [
            "--profile", str(profile_5_yaml),
            "--offline",
            "--output-dir", str(out_dir),
            "--format", "both",
            "--quiet",
        ]
        exit_code = main(argv)
        assert exit_code == 0

        # Verify dedicated tracker was generated for Profile 5 without any code edits
        xlsx_5 = out_dir / "profile_5_especialista_geotecnia_vial_tracker.xlsx"
        csv_5 = out_dir / "profile_5_especialista_geotecnia_vial_tracker.csv"
        assert xlsx_5.exists()
        assert csv_5.exists()

        loaded_5 = ExcelTracker.load(xlsx_5)
        assert isinstance(loaded_5, list)


# =============================================================================
# E2E TEST 5: VALIDATE ALL ACCEPTANCE CRITERIA FROM ORIGINAL_REQUEST.MD
# =============================================================================

class TestE2EAcceptanceCriteriaVerification:
    """Explicitly verifies every single Acceptance Criterion from ORIGINAL_REQUEST.md lines 36-52."""

    def test_criterion_1_functional_extraction_and_data_normalization(self, canonical_benchmark_vacancies):
        """AC 1 (Line 39): The extraction engine recovers functional vacancies with valid URLs and
        recent dates from recognized sources in Colombia/remoto.
        """
        engine = JobAggregationEngine(custom_scrapers={
            "mock_portal": MockCustomScraper(canonical_benchmark_vacancies, name="mock_portal")
        })

        results = engine.run(queries=["ingeniero civil"], portals=["mock_portal"])
        assert len(results) == 9

        for vac in results:
            assert vac.id and len(vac.id) == 16, f"Invalid dedup hash: {vac.id}"
            assert vac.title, "Vacancy missing title"
            assert vac.company, "Vacancy missing company"
            assert vac.direct_url.startswith("http"), f"Invalid direct URL: {vac.direct_url}"
            assert vac.modality in ("remoto", "presencial", "hibrido", "desconocido")
            assert vac.publication_date, "Missing publication date"

    def test_criterion_2_network_error_resilience_and_fault_isolation(self):
        """AC 2 (Line 40): System handles network errors and page structure changes
        resiliently without terminating abruptly.
        """
        # Create a failing scraper simulating HTTP 500 / Network Error
        failing_scraper = MagicMock()
        failing_scraper.search.side_effect = ConnectionError("Network unreachable / Timeout")

        # Create a healthy scraper
        healthy_vacancy = NormalizedVacancy(
            id="healthy_vac_0001",
            title="Ingeniera Civil Vías Remoto",
            company="Constructora Resiliente",
            direct_url="https://portal.com/job/healthy",
            full_description="100% teletrabajo en Colombia.",
            modality="remoto",
            source_portal="computrabajo",
        )
        healthy_scraper = MockCustomScraper([healthy_vacancy], name="computrabajo")

        engine = JobAggregationEngine(custom_scrapers={
            "failing_portal": failing_scraper,
            "computrabajo": healthy_scraper,
        })

        # Run aggregation across both portals: failing portal must NOT crash the engine
        vacancies = engine.run(queries=["ingeniero"], portals=["failing_portal", "computrabajo"])

        assert len(vacancies) == 1
        assert vacancies[0].id == "healthy_vac_0001"
        assert "failing_portal" in engine.last_run_stats["failures"]

    def test_criterion_3_presential_obra_vacancies_filtered_or_penalized(self, canonical_benchmark_vacancies):
        """AC 3 (Line 43): 100% on-site obra or non-virtual vacancies are filtered or penalized."""
        profile = load_profile("profile_1_civil_vias_sst")
        verifier = RemoteVerifier()
        scorer = AffinityScorer()

        presential_jobs = [v for v in canonical_benchmark_vacancies if v.extra_metadata.get("category") == "presential"]
        assert len(presential_jobs) == 3

        for vac in presential_jobs:
            mod_eval = verifier.evaluate(vac, profile)
            # Must be marked as disqualified or not strictly remote
            assert mod_eval.is_disqualified or not mod_eval.is_strictly_remote
            assert mod_eval.modality_score == 0.0

            breakdown = scorer.evaluate(vac, profile)
            # Disqualified or heavily penalized score <= 20
            assert breakdown.total_score <= 20.0
            assert breakdown.is_disqualified is True or breakdown.penalties > 0

    def test_criterion_4_score_greater_than_70_matches_civil_vias_sst(self, canonical_benchmark_vacancies):
        """AC 4 (Line 44): Vacancies with affinity score > 70 authentically match civil engineering,
        vías/infraestructura, and SST profiles.
        """
        profile = load_profile("profile_1_civil_vias_sst")
        scorer = AffinityScorer()

        ideal_jobs = [v for v in canonical_benchmark_vacancies if v.extra_metadata.get("category") == "ideal"]
        nonrelevant_jobs = [v for v in canonical_benchmark_vacancies if v.extra_metadata.get("category") == "nonrelevant"]
        assert len(ideal_jobs) == 3
        assert len(nonrelevant_jobs) == 3

        # Ideal jobs must score >= 70
        for vac in ideal_jobs:

            bd = scorer.evaluate(vac, profile)
            assert bd.total_score >= 70.0, f"Ideal job {vac.id} scored {bd.total_score} < 70"
            assert not bd.is_disqualified
            assert "COPNIA" in bd.explanatory_summary or "SST" in bd.explanatory_summary or "vías" in bd.explanatory_summary.lower()

        # Non-relevant jobs must score < 50
        for vac in nonrelevant_jobs:
            bd = scorer.evaluate(vac, profile)
            assert bd.total_score < 50.0, f"Non-relevant job {vac.id} scored {bd.total_score} >= 50"

    def test_criterion_5_and_6_storage_formatting_and_state_preservation(self, tmp_path: Path):
        """AC 5 & 6 (Lines 50-52): Automated pytest suite validates extraction, dedup,
        score filtering, Excel/CSV persistence, legible format, clickable links, and status preservation.
        """
        excel_out = tmp_path / "ac_test.xlsx"
        csv_out = tmp_path / "ac_test.csv"

        entry = TrackerEntry(
            id="ac_test_0001",
            title="Directora Vial SST",
            company="Empresa AC",
            direct_url="https://portal.com/job/ac1",
            full_description="Descripción de prueba",
            source_portal="computrabajo",
            score=92.0,
            score_percentage="92%",
            estado="Por revisar",
            notas_usuario="Revisar urgente",
        )

        ExcelTracker.save([entry], excel_out, preserve_existing=False)
        CSVTracker.save([entry], csv_out, preserve_existing=False)

        # Excel verified
        wb = openpyxl.load_workbook(excel_out)
        ws = wb.active
        assert ws.cell(row=2, column=10).value == "Por revisar"
        assert ws.cell(row=2, column=11).value == "Revisar urgente"
        assert "=HYPERLINK" in str(ws.cell(row=2, column=7).value)
        wb.close()

        # CSV verified
        raw_csv = csv_out.read_bytes()
        assert raw_csv.startswith(b"\xef\xbb\xbf")  # UTF-8 BOM


# =============================================================================
# E2E TEST 6: CHALLENGER M4 HARDENING MITIGATIONS VERIFICATION
# =============================================================================

class TestE2EChallengerM4HardeningMitigations:
    """Explicitly verify the 5 adversarial hardening mitigations identified by Challenger M4."""

    def test_mitigation_1_excel_control_characters_cleaned(self, tmp_path: Path):
        """Mitigation 1: Control characters in strings do not crash openpyxl with IllegalCharacterError."""
        excel_file = tmp_path / "cleaned_ctrl.xlsx"
        entry_with_ctrl = TrackerEntry(
            id="ctrl_clean_1",
            title="Título con form feed \x0c y escape \x1b",
            company="Empresa \x07 bell",
            direct_url="https://portal.com/job/\x0btest",
            full_description="Detalle con caracteres de control \x00\x08\x0e\x1f válidamente limpiados",
            source_portal="test",
            score=85.0,
            explanatory_summary="Resumen con \x0c control",
            notas_usuario="Notas con \x1b control",
        )

        # Must succeed without IllegalCharacterError
        ExcelTracker.save([entry_with_ctrl], excel_file, preserve_existing=False)
        assert excel_file.exists()

        loaded = ExcelTracker.load(excel_file)
        assert len(loaded) == 1
        assert "\x0c" not in loaded[0].title
        assert "\x07" not in loaded[0].company

    def test_mitigation_2_corrupt_and_zero_byte_excel_files_handled_gracefully(self, tmp_path: Path):
        """Mitigation 2: 0-byte or corrupted xlsx files return empty list instead of BadZipFile."""
        zero_byte = tmp_path / "empty.xlsx"
        zero_byte.write_bytes(b"")

        corrupt = tmp_path / "corrupt.xlsx"
        corrupt.write_bytes(b"PK\x03\x04corrupted fake excel stream")

        # ExcelTracker.load returns []
        assert ExcelTracker.load(zero_byte) == []
        assert ExcelTracker.load(corrupt) == []

        # IncrementalTracker.merge on corrupt file returns [] or freshly initialized entries without crashing
        fresh = [TrackerEntry(id="fresh_after_corrupt", title="Job 1", company="C", direct_url="https://x.com", full_description="D", source_portal="t", score=80.0)]
        merged = IncrementalTracker.merge(corrupt, fresh)
        assert len(merged) == 1
        assert merged[0].id == "fresh_after_corrupt"

    def test_mitigation_3_first_run_duplicate_id_deduplication(self, tmp_path: Path):
        """Mitigation 3: IncrementalTracker.merge() deduplicates fresh_entries by ID on first run."""
        non_existent = tmp_path / "does_not_exist.xlsx"
        duplicate_entries = [
            TrackerEntry(id="dup_same_id", title="Job Version 1", company="C", direct_url="https://x.com", full_description="D", source_portal="t", score=60.0),
            TrackerEntry(id="dup_same_id", title="Job Version 2 Updated", company="C", direct_url="https://x.com", full_description="D", source_portal="t", score=88.0),
        ]

        merged = IncrementalTracker.merge(non_existent, duplicate_entries)
        # Must deduplicate to exactly 1 entry with highest score
        assert len(merged) == 1
        assert merged[0].id == "dup_same_id"
        assert merged[0].score == 88.0

    def test_mitigation_4_score_bounds_clamped_on_load(self, tmp_path: Path):
        """Mitigation 4: Scores exceeding 100% in external CSV or Excel are clamped to 100.0 without ValidationError."""
        csv_file = tmp_path / "score_over_100.csv"
        csv_file.write_text(
            "ID/Hash,Fecha Detección,Título,Empresa,Score %,Modalidad,Enlace Directo,Salario,Resumen de Afinidad,Estado,Notas Usuario\n"
            "over100,2026-09-20,Ingeniero,Empresa,135%,remoto,https://portal.com,,summary,Nueva,\n",
            encoding="utf-8-sig",
        )

        # Must load cleanly with clamped score 100.0
        loaded_csv = CSVTracker.load(csv_file)
        assert len(loaded_csv) == 1
        assert loaded_csv[0].score == 100.0

    def test_mitigation_5_date_detection_header_regex_prioritization(self, tmp_path: Path):
        """Mitigation 5: 'Fecha Detección' is matched accurately without colliding with 'Fecha Postulación'."""
        excel_file = tmp_path / "date_headers.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append([
            "ID/Hash", "Fecha Detección", "Título", "Empresa", "Score %",
            "Modalidad", "Enlace Directo", "Salario", "Resumen de Afinidad",
            "Estado", "Notas Usuario", "Fecha Postulación"
        ])
        ws.append([
            "date_test_1", "2026-09-01 08:30", "Ingeniera Civil", "Empresa X", 0.85,
            "Remoto", "https://x.com", "$10M", "Summary", "Nueva", "Nota", "2026-09-20"
        ])
        wb.save(excel_file)
        wb.close()

        loaded = ExcelTracker.load(excel_file)
        assert len(loaded) == 1
        # Must match 2026-09-01 08:30, NOT 2026-09-20
        assert loaded[0].fecha_deteccion == "2026-09-01 08:30"
